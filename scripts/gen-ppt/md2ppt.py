#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
md2ppt.py —— Markdown → python-pptx 脚本 → PPTX

用法:
    python md2ppt.py <deck.md> [out.pptx]

Markdown 约定:
    ---            YAML 元数据（title/core/platform/date）→ 封面
    # 标题         新幻灯片
    [图:name fs=16] 插入预定义矢量图（可指定字号）
    - 要点         列表（缩进 2 空格 = 二级）
    | a | b |      Markdown 表格
    > 文本         脚注；`> [24] 文本` 指定字号并加粗
    <!-- tbl size=16 colw=3,2,2 -->     表格属性（字号/列宽）
    <!-- bullets size=24 spacing=1.5 --> 列表属性（字号/行距）
    支持 **粗体** / *斜体*
"""
import re
import sys
import os

REPO_DIR = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '../..'))

# =====================================================================
# 图函数库（fs 参数可调；图占上半页）
# =====================================================================
FIGURES = {}

FIGURES['roadmap'] = '''
def draw_roadmap(s, fs=14):
    ms = [('M1', '上板闭环', '从零到上板闭环', GREEN),
          ('M2', '调试流程与环境', '流程/环境/规范落地', GREEN),
          ('M3', '完整 Core + 测试向量', 'TL-UL 主干 + RVV + 606 用例', GREEN),
          ('M4', '存储系统扩展', 'TCM 扩容 / SPI / DDR', ORANGE),
          ('M5', '矩阵计算 + benchmark', '矩阵运算与基准评测', GRAY)]
    for i, (tag, name, desc, col) in enumerate(ms):
        x = 0.55 + i * 2.5
        box(s, tag + '\\n' + name, x, 1.45, 2.3, 0.9, fill=col, tc=WHITE, fs=fs, bold=True)
        box(s, desc, x, 2.45, 2.3, 0.75, fs=fs)
        if i < 4:
            arrow(s, x + 2.35, 1.75, 0.1, h=0.14)
'''

FIGURES['storage'] = '''
def draw_storage(s, fs=16):
    box(s, '核（SCore + RvvCore）', 4.6, 1.25, 4.1, 0.55, fill=DARK, tc=WHITE, fs=fs, bold=True)
    box(s, '取指：UncachedFetch（无 Cache）', 3.4, 1.95, 6.5, 0.55, fill=RGBColor(0xF5,0xE0,0xD6), fs=fs)
    box(s, 'TCM：ITCM（8K~1M）+ DTCM（32K~1M）· 单周期 · FabricArbiter', 2.1, 2.65, 9.1, 0.55, fill=RGBColor(0xBF,0xD7,0xEE), fs=fs)
    box(s, '外部：SRAM · ROM · DDR（规划）', 0.8, 3.35, 11.7, 0.55, fill=RGBColor(0xA8,0xCB,0xE8), fs=fs)
    for y in (1.85, 2.55, 3.25):
        arrow(s, 6.55, y, 0.0, 0.1, shape=MSO_SHAPE.DOWN_ARROW)
'''

FIGURES['bus'] = '''
def draw_bus(s, fs=16):
    hosts = ['coralnpu_core', 'uart_host', 'spi2tlul（规划）', 'dma（规划）']
    for i, h in enumerate(hosts):
        box(s, h, 0.7 + i * 3.15, 1.35, 2.6, 0.55, fs=fs)
    box(s, 'CoralNPUXbar —— TileLink-UL 交叉开关', 0.7, 2.05, 11.9, 0.6, fill=DARK, tc=WHITE, fs=fs, bold=True)
    devs = ['coralnpu_device', 'sram', 'clint / plic', 'gpio', 'rom', 'spi / dma']
    for i, d in enumerate(devs):
        box(s, d, 0.7 + i * 2.05, 2.8, 1.8, 0.55, fs=fs)
'''

FIGURES['swstack'] = '''
def draw_swstack(s, fs=18):
    sw = [('工具链', 'RISC-V clang（host_clang）· crt0 · TCM 链接脚本'),
          ('运行库', 'litert-micro（conv / depthwise_conv / fully_connected）· rvv_opt.h'),
          ('示例', 'hello_world · rvv_add_intrinsic · MobileNet V1'),
          ('验证', 'cocotb · npusim · systemc · uvm · vcs_sim · verilator_sim')]
    for i, (name, desc) in enumerate(sw):
        y = 1.3 + i * 0.78
        box(s, name, 0.5, y, 1.7, 0.7, fill=DARK, tc=WHITE, fs=fs, bold=True)
        box(s, desc, 2.35, y, 10.5, 0.7, fs=fs)
'''

FIGURES['trim'] = '''
def draw_trim(s, fs=18):
    box(s, '上游 chip_nexus（完整 SoC 设计）', 0.6, 1.35, 12.1, 0.6, fill=GRAY, tc=WHITE, fs=fs, bold=True)
    arrow(s, 6.6, 2.02, 0.0, 0.32, shape=MSO_SHAPE.DOWN_ARROW, color=RED)
    box(s, '裁剪后（我们，M3）：CoreTlul + CoralNPUXbar + 外设 + uart_host', 0.6, 2.45, 12.1, 0.6, fill=GREEN, tc=WHITE, fs=fs, bold=True)
    box(s, '删除：ISP / DDR / spi2tlul / dma / spi_master(_flash)', 0.6, 3.2, 12.1, 0.55, fs=fs)
'''

FIGURES['loadpath'] = '''
def draw_loadpath(s, fs=14):
    chain = [('PC 串口', 0.7), ('UART\\n收发', 2.75), ('host_cmd_fsm\\n命令解析', 4.8),
             ('Axi2TLUL\\n桥', 6.85), ('Xbar\\nuart_host', 8.9), ('核 TCM/CSR\\n· SRAM · 外设', 10.95)]
    for t, x in chain:
        box(s, t, x, 1.6, 1.7, 0.95, fs=fs)
        if x < 10.9:
            arrow(s, x + 1.78, 1.95, 0.1)
'''

FIGURES['flow'] = '''
def draw_flow(s, fs=18):
    box(s, '100MHz 差分', 0.7, 1.35, 2.6, 0.6, fs=fs)
    arrow(s, 3.35, 1.57, 0.4)
    box(s, 'MMCM ×12/60', 3.85, 1.35, 2.6, 0.6, fill=RGBColor(0xD6,0xE4,0xF5), fs=fs)
    arrow(s, 6.5, 1.57, 0.4)
    box(s, '20MHz 核时钟（单时钟域）', 7.0, 1.35, 5.6, 0.6, fill=RGBColor(0xD5,0xEF,0xDC), fs=fs)
    flow = [('Chisel RTL', 0.7), ('bazel 生成 SV', 3.3), ('Vivado 综合', 5.9),
            ('bitstream', 8.5), ('上板', 10.9)]
    for i, (t, x) in enumerate(flow):
        box(s, t, x, 2.3, 2.3, 0.8, fs=fs)
        if i < 4:
            arrow(s, x + 2.35, 2.62, 0.12)
'''

FIGURES['soc_arch'] = '''
def draw_soc_arch(s, fs=16):
    # 核心 rvv_core（居中）
    box(s, 'rvv_core (CoreTlul)\\n标量 + RVV + FPU + LSU', 4.3, 1.25, 4.9, 1.0,
        fill=DARK, tc=WHITE, fs=fs, bold=True)
    # 下方：存储层次（类似存储体系页）
    box(s, 'TCM：ITCM + DTCM（单周期）', 4.3, 2.4, 4.9, 0.55, fill=RGBColor(0xBF,0xD7,0xEE), fs=fs)
    box(s, '外部：SRAM / ROM / DDR（规划）', 4.3, 3.1, 4.9, 0.55, fill=RGBColor(0xA8,0xCB,0xE8), fs=fs)
    arrow(s, 6.75, 2.3, 0.0, 0.1, shape=MSO_SHAPE.DOWN_ARROW)
    # 右侧：外设
    box(s, '外设', 9.7, 1.25, 3.0, 0.5, fill=LIGHT, fs=fs, bold=True)
    box(s, 'sram', 9.7, 1.85, 3.0, 0.5, fs=fs)
    box(s, 'clint / plic', 9.7, 2.45, 3.0, 0.5, fs=fs)
    box(s, 'gpio / rom', 9.7, 3.05, 3.0, 0.5, fs=fs)
    # 左侧：加载通路
    box(s, 'PC → UART\\nhost_cmd_fsm / Axi2TLUL', 0.6, 1.6, 3.3, 0.9, fs=fs)
    # 连接（Xbar 弱化）
    box(s, 'CoralNPUXbar（TileLink-UL 连接）', 4.3, 3.8, 4.9, 0.5, fill=LIGHT, fs=fs)
'''

FIGURES['official_block'] = '''
def draw_official_block(s, fs=14):
    img = os.path.join(r"__REPO__", 'coralnpu/doc/images/arch_data_flow.png')
    s.shapes.add_picture(img, Inches(3.1), Inches(1.3), height=Inches(4.0))
'''

# =====================================================================
# 各图底部 y 坐标
# =====================================================================
FIG_BOTTOM = {
    'roadmap': 3.3, 'storage': 4.0, 'bus': 3.5, 'swstack': 4.45,
    'trim': 3.85, 'loadpath': 2.7, 'flow': 3.3, 'soc_arch': 4.4,
    'official_block': 5.35,
}

# =====================================================================
# 生成脚本的头
# =====================================================================
HEADER = '''#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""由 md2ppt.py 从 Markdown 自动生成 —— 请勿手工修改，改 Markdown 后重跑 md2ppt.py"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

DARK   = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x2E, 0x75, 0xB6)
LIGHT  = RGBColor(0xE8, 0xF0, 0xF8)
GRAY   = RGBColor(0x59, 0x59, 0x59)
RED    = RGBColor(0xC0, 0x39, 0x2B)
GREEN  = RGBColor(0x27, 0x8A, 0x4A)
ORANGE = RGBColor(0xD9, 0x7A, 0x06)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
FONT   = 'Microsoft YaHei'

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
PAGE = [0]

def slide(title=None):
    s = prs.slides.add_slide(BLANK)
    if title:
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.28), Inches(12.3), Inches(0.75))
        p = tb.text_frame.paragraphs[0]
        r = p.add_run(); r.text = title
        r.font.size = Pt(25); r.font.bold = True; r.font.color.rgb = DARK; r.font.name = FONT
        ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.02), Inches(12.3), Pt(2.2))
        ln.fill.solid(); ln.fill.fore_color.rgb = ACCENT; ln.line.fill.background()
    return s

def add_runs(p, text, size, color, bold_default=False, font=None):
    import re
    f = font or FONT
    for part in re.split(r'(\\*\\*[^*]+\\*\\*|\\*[^*]+\\*)', text):
        if not part:
            continue
        r = p.add_run()
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            r.text = part[2:-2]; r.font.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r.text = part[1:-1]; r.font.italic = True
        else:
            r.text = part; r.font.bold = bold_default
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = f

def bullets(s, items, left=0.6, top=1.25, width=12.1, height=None, size=18, spacing=None):
    if height is None:
        height = (size * 0.021) * len(items) + 0.25
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        text, lvl = (it if isinstance(it, tuple) else (it, 0))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = lvl; p.space_after = Pt(3)
        if spacing:
            p.line_spacing = spacing
        add_runs(p, text, size - lvl * 1.5, DARK if lvl == 0 else GRAY)
    return tb

def quote(s, text, top, size=14):
    tb = s.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12.1), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    add_runs(p, text, size, GRAY)
    return tb

def para(s, text, top, size=20):
    tb = s.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12.1), Inches(0.45))
    p = tb.text_frame.paragraphs[0]
    add_runs(p, text, size, DARK, bold_default=True)
    return tb

def table(s, headers, rows, left=0.6, top=1.25, width=12.1, height=5.0, fs=14, colw=None):
    shp = s.shapes.add_table(len(rows) + 1, len(headers), Inches(left), Inches(top),
                             Inches(width), Inches(height))
    tbl = shp.table
    if colw:
        for j, w in enumerate(colw):
            if j < len(headers):
                tbl.columns[j].width = Inches(w)
    row_h = height / (len(rows) + 1)
    for r in tbl.rows:
        r.height = Inches(row_h)
    for j, h in enumerate(headers):
        c = tbl.cell(0, j); c.text = ''
        add_runs(c.text_frame.paragraphs[0], str(h), fs, WHITE, bold_default=True)
        c.fill.solid(); c.fill.fore_color.rgb = ACCENT
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = tbl.cell(i + 1, j); c.text = ''
            add_runs(c.text_frame.paragraphs[0], str(v), fs, DARK)
            c.fill.solid(); c.fill.fore_color.rgb = LIGHT if i % 2 == 0 else WHITE
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
    return tbl

def box(s, text, x, y, w, h, fill=LIGHT, tc=DARK, fs=14, bold=False,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=ACCENT):
    shp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_runs(p, text, fs, tc, bold_default=bold)
    return shp

def arrow(s, x, y, w, h=0.16, color=ACCENT, shape=MSO_SHAPE.RIGHT_ARROW):
    shp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = color; shp.line.fill.background()
    return shp

def footer(s, text):
    PAGE[0] += 1
    tb = s.shapes.add_textbox(Inches(0.5), Inches(7.02), Inches(12.3), Inches(0.35))
    p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = text + '   |   ' + str(PAGE[0])
    r.font.size = Pt(14); r.font.color.rgb = GRAY; r.font.name = FONT

FOOT = 'coralnpu RISC-V NPU 上板验证与性能评估'
'''

FOOTER = '''
out = r"__OUT__"
prs.save(out)
print('已生成:', out, '页数:', len(prs.slides._sldIdLst))
'''


def parse_md(text):
    m = re.match(r'^---\n(.*?)\n---\n', text, re.S)
    yaml = {}
    if m:
        for line in m.group(1).splitlines():
            if ':' in line:
                k, v = line.split(':', 1)
                yaml[k.strip()] = v.strip().strip('"')
        text = text[m.end():]
    pages = []
    cur = None
    lines = text.splitlines()
    i = 0
    tbl_attr = {}
    bl_attr = {}
    q_attr = {}
    while i < len(lines):
        ln = lines[i]
        if ln.startswith('# '):
            if cur:
                pages.append(cur)
            cur = {'title': ln[2:].strip(), 'blocks': []}
            i += 1
            continue
        if cur is None:
            i += 1
            continue
        ma = re.match(r'^<!--\s*(tbl|bullets|quote)\s+(.*?)\s*-->', ln.strip())
        if ma:
            attrs = {}
            for kv in ma.group(2).split():
                if '=' in kv:
                    k, v = kv.split('=', 1); attrs[k] = v
            if ma.group(1) == 'tbl':
                tbl_attr = attrs
            elif ma.group(1) == 'bullets':
                bl_attr = attrs
            else:
                q_attr = attrs
            i += 1
            continue
        mf = re.match(r'^\[图:([a-z_]+)(?:\s+fs=(\d+))?\]', ln.strip())
        if mf:
            cur['blocks'].append(('fig', mf.group(1), int(mf.group(2)) if mf.group(2) else None))
            i += 1
            continue
        if ln.strip().startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                if not re.match(r'^[\s\-:|]+$', lines[i].strip().strip('|')):
                    rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            if rows:
                cur['blocks'].append(('table', rows[0], rows[1:], tbl_attr))
                tbl_attr = {}
            continue
        if re.match(r'^\s*-\s+', ln):
            items = []
            while i < len(lines) and re.match(r'^\s*-\s+', lines[i]):
                indent = len(lines[i]) - len(lines[i].lstrip())
                items.append((re.sub(r'^\s*-\s+', '', lines[i]).strip(), 1 if indent >= 2 else 0))
                i += 1
            cur['blocks'].append(('bullets', items, bl_attr))
            bl_attr = {}
            continue
        if ln.strip().startswith('>'):
            cur['blocks'].append(('quote', ln.strip().lstrip('>').strip(), q_attr))
            q_attr = {}
            i += 1
            continue
        if ln.strip():
            cur['blocks'].append(('para', ln.strip()))
            i += 1
            continue
        i += 1
    if cur:
        pages.append(cur)
    return yaml, pages


def py_str(s):
    return repr(s)


def gen_py(yaml, pages, out_py, out_pptx_name):
    code = [HEADER]
    for name in sorted(set(b[1] for p in pages for b in p['blocks'] if b[0] == 'fig')):
        fig = FIGURES.get(name, f'def draw_{name}(s, fs=14):\n    pass\n')
        code.append(fig.replace('__REPO__', REPO_DIR))
    code.append('\n# ===== 封面 =====')
    title = yaml.get('title', '')
    core = yaml.get('core', yaml.get('subtitle', ''))
    platform = yaml.get('platform', '')
    date = yaml.get('date', '')
    code.append(f'''
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = DARK; bg.line.fill.background()
tb = s.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(1.6))
p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = {py_str(title)}
r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
tb2 = s.shapes.add_textbox(Inches(0.8), Inches(4.1), Inches(11.7), Inches(2.0))
tf2 = tb2.text_frame
for i, t in enumerate([{py_str(core)}, {py_str(platform)}, '', '', {py_str(date)}]):
    p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
    p.alignment = PP_ALIGN.CENTER; p.space_after = Pt(6)
    r = p.add_run(); r.text = t; r.font.size = Pt(18); r.font.color.rgb = RGBColor(0xC8,0xD8,0xE8); r.font.name = FONT
''')
    for p in pages:
        title = p['title']
        code.append(f'\n# ===== {title} =====')
        code.append(f's = slide({py_str(title)})')
        y = 1.25
        for blk in p['blocks']:
            kind = blk[0]
            if kind == 'fig':
                fs_arg = f', fs={blk[2]}' if blk[2] else ''
                code.append(f'draw_{blk[1]}(s{fs_arg})')
                y = FIG_BOTTOM.get(blk[1], 3.7) + 0.15
            elif kind == 'bullets':
                items = blk[1]
                attrs = blk[2] if len(blk) > 2 else {}
                n = len(items)
                size = int(attrs.get('size', 20 if n <= 4 else 18))
                spacing = attrs.get('spacing')
                top_v = float(attrs['top']) if 'top' in attrs else y
                code.append('bullets(s, [')
                for t, lvl in items:
                    code.append(f'    ({py_str(t)}, {lvl}),')
                code.append(f'], top={top_v:.2f}, size={size}' + (f', spacing={spacing}' if spacing else '') + ')')
                y = top_v + (size * 0.021 * (float(spacing) if spacing else 1.0)) * n + 0.30
            elif kind == 'table':
                rows_n = len(blk[2])
                attrs = blk[3] if len(blk) > 3 else {}
                fs = int(attrs.get('size', 14))
                colw = [float(x) for x in attrs['colw'].split(',')] if 'colw' in attrs else None
                h = min(0.42 * (rows_n + 1) + 0.2, 5.4)
                code.append(f'table(s, {py_str(blk[1])}, {py_str(blk[2])}, top={y:.2f}, height={h:.2f}, fs={fs}, colw={colw})')
                y += h + 0.15
            elif kind == 'para':
                code.append(f'para(s, {py_str(blk[1])}, {y:.2f})')
                y += 0.45
            elif kind == 'quote':
                txt = blk[1]
                qa = blk[2] if len(blk) > 2 else {}
                qtop = float(qa['top']) if 'top' in qa else y
                mq = re.match(r'^\[(\d+)\]\s*(.*)$', txt)
                if mq:
                    qsize = int(mq.group(1)); qtext = mq.group(2)
                    code.append(f'quote(s, {py_str("**" + qtext + "**")}, {qtop:.2f}, size={qsize})')
                    y = qtop + 0.6
                else:
                    # 佐证/说明：统一固定到页面底部（靠下）
                    code.append(f'quote(s, {py_str(txt)}, 6.45, size=14)')
        code.append(f'footer(s, FOOT)')
    code.append(FOOTER.replace('__OUT__', out_pptx_name))
    with open(out_py, 'w') as f:
        f.write('\n'.join(code))
    print('已生成脚本:', out_py)


def main():
    if len(sys.argv) < 2:
        print('用法: python md2ppt.py <deck.md> [out.pptx]')
        sys.exit(1)
    md = sys.argv[1]
    out_pptx = os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 else os.path.abspath('report.pptx')
    out_py = os.path.join(os.path.dirname(out_pptx), 'gen.py')
    yaml, pages = parse_md(open(md, encoding='utf-8').read())
    gen_py(yaml, pages, out_py, out_pptx)


if __name__ == '__main__':
    main()
